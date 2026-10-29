"""Build the detailed formal operations manual; render and inspect before publication."""
import argparse
import json
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from enterprise_manual_content import CHAPTERS

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--screenshots', type=Path, required=True)
parser.add_argument('--output', type=Path, required=True)
parser.add_argument('--url', default='http://192.168.1.11:18860')
parser.add_argument('--version', required=True)
parser.add_argument('--contents-pages', type=Path, help='Verified heading-to-page mapping from the rendered PDF')
args = parser.parse_args()
pages = json.loads(args.contents_pages.read_text()) if args.contents_pages else {}
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(3.7), Cm(3.5)
s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)
s.footer_distance = Cm(2.5)
for name, font, size, line in [('Normal','仿宋',16,28),('Title','方正小标宋简体',22,32),('Heading 1','黑体',16,28),('Heading 2','楷体',16,28),('Caption','宋体',12,20),('Body Text','仿宋',14,24)]:
    style=doc.styles[name]
    style.font.name=font
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.bold=False
    style.paragraph_format.line_spacing=Pt(line)
    style.paragraph_format.space_before=Pt(0)
    style.paragraph_format.space_after=Pt(0)
    style.paragraph_format.first_line_indent=Pt(32 if name=='Normal' else 0)
    style.paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
footer=s.footer.paragraphs[0];footer.alignment=1;footer.style=doc.styles['Caption']
footer.add_run('— ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' —')
doc.core_properties.title='企业AI智能运维平台操作手册'
doc.core_properties.subject='功能说明 操作步骤 结果核验 常见问题'
doc.core_properties.author=''
doc.core_properties.keywords='运维 监控 操作手册'
NUMBERS=['一','二','三','四','五','六','七','八','九','十','十一','十二','十三','十四','十五','十六','十七','十八','十九','二十']

def para(text):
    return doc.add_paragraph(text)

def new_heading(text):
    p=doc.add_heading(text,1)
    p.paragraph_format.page_break_before=True
    return p

def figure(name, caption, height=8):
    path=args.screenshots/f'{name}.png'
    with Image.open(path) as im:
        width=min(15.6,height*im.width/im.height)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=0
    p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True;p.alignment=1
    p.add_run().add_picture(str(path),width=Cm(width))
    # Alt text makes the figure identifiable in screen readers and Word.
    for node in p._p.iter(qn('wp:docPr')): node.set('descr',caption)
    p=doc.add_paragraph(caption,'Caption');p.alignment=1

def field_table(fields):
    table=doc.add_table(rows=0,cols=2)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    table.columns[0].width=Cm(3.3);table.columns[1].width=Cm(12.3)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement(f'w:{side}');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    for index,(label,value) in enumerate([('字段或区域','含义与使用范围')]+[tuple(f.split('：',1)) for f in fields]):
        row=table.add_row()
        # Keep each short field definition intact; never use fixed row heights.
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        if index==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for col,(cell,text) in enumerate(zip(row.cells,[label,value])):
            cell.width=Cm(3.3 if col==0 else 12.3);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text
            pr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                item=OxmlElement(f'w:{side}');item.set(qn('w:w'),'80');item.set(qn('w:type'),'dxa');margin.append(item)
            pr.append(margin)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=0;p.paragraph_format.line_spacing=Pt(20)
                for run in p.runs:
                    run.font.size=Pt(13);run.font.name='仿宋'
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'仿宋')
    doc.add_paragraph().paragraph_format.line_spacing=Pt(6)

p=doc.add_paragraph('企业AI智能运维平台\n操作手册','Title');p.alignment=1
for text in ['详细操作版',f'适用系统版本 {args.version}',f'访问地址 {args.url}','2026年10月']:
    p=para(text);p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.space_before=Pt(28)
# Fixed two-page contents prevents page-number updates from shifting body pagination.
for start in [0,10]:
    new_heading('目　录' if start==0 else '目　录　续')
    entries=([('使用说明','使用说明')] if start==0 else [])+[(f'{NUMBERS[i]}、{c["title"]}',f'{NUMBERS[i]}、{c["title"]}') for i,c in enumerate(CHAPTERS) if start<=i<start+10]
    if start==10:
        for title in ['附录　操作核验清单','附录　阈值保存结果','附录　告警关闭结果','附录　流程完成结果','附录　知识编辑填写','附录　页面入口图示']:
            entries.append((title,title))
    for title,key in entries:
        p=para(f'{title}　{pages.get(key,"待核定")}');p.paragraph_format.first_line_indent=0;p.paragraph_format.space_after=Pt(8)
new_heading('使用说明')
for text in [
 '本手册面向使用企业AI智能运维平台的运维操作人员，说明登录、监控检查、告警处置、流程操作和知识维护的方法。按章节完成操作后，应核对页面结果，不只确认按钮已经点击。',
 '全书按“功能与字段—操作步骤—结果核验”组织。截图采用中文浅色界面，重点展示列表筛选、详情窗口与确认步骤；不同浏览器的字体、时间格式和校验提示可能略有差异。',
 '系统使用本地样本快照、模拟指标和预设分析。操作数据保存在当前浏览器，不跨电脑同步。登录仅检查本地输入形式，流程不连接实际设备，阈值不下发到采集端。',
 '首次使用建议依次完成登录、总览浏览、主机详情检查、告警确认、流程执行、执行记录核对和知识新增。需要保留的文本，应在恢复数据或清理浏览器前另存副本。',
 '文中对象值与数量均以初始样本为基准。部分截图展示操作后的状态，例如阈值改为 88、告警已关闭或新增执行记录；这些变化不代表真实生产系统发生变更。',
 '截图中的账号仅为输入示例，密码为临时输入；平台不存在统一固定口令。图片中的数据时间取自独立截图会话，不作为本手册的发布日期。',

]:para(text)
for i,c in enumerate(CHAPTERS):
    title=f'{NUMBERS[i]}、{c["title"]}'
    new_heading(title);para(c['intro'])
    doc.add_heading('（一）功能与字段',2);field_table(c['fields'])
    doc.add_heading('（二）操作步骤',2)
    for step,text in enumerate(c['steps'],1):para(f'{step}. {text}')
    new_heading(f'{NUMBERS[i]}、{c["title"]}　操作图示')
    figure('knowledge-saved' if c['detail']=='knowledge-new' else c['detail'],f'图 {i+1}－1　{"新增知识保存后的详情" if c["detail"]=="knowledge-new" else c["captions"][1]}',height=11)
    doc.add_heading('（三）结果核验与注意事项',2)
    for text in c['checks']:
        p=doc.add_paragraph(text,'Body Text');p.paragraph_format.first_line_indent=Pt(28)
new_heading('附录　操作核验清单')
for title,items in [
 ('（一）一次主机高负载检查',[
 '在主机监控搜索 order-api-01，打开近 6 小时详情，核对名称、地址和 CPU 指标。',
 '通过资源台账查看 res-1 与 INC-1201，再到告警详情确认状态。若事件已关闭，不要期待确认按钮仍可点击。',
 '在智能助手查看高 CPU 场景，阅读证据和知识；打开关联流程并确认一次本地执行。',
 '核对执行列表新增记录，再新建巡检知识，填写对象、检查内容与结论。流程完成不代表真实故障已经处理。']),
 ('（二）一次规则修改与复查',[
 '记录 CPU 规则原值 85，将其改为 88 并保存。刷新后核对新值仍存在。',
 '确认既有告警状态没有自动改变；这是规则与样本快照相互独立的结果。',
 '需要还原时，将 CPU 规则重新改为 85。不要通过恢复业务数据来尝试恢复规则。']),
 ('（三）反馈问题时提供的信息',[
 '记录访问地址、运行版本、页面名称、发生时间、对象编号及重现步骤。保留错误提示与截图，说明是否刷新或清理站点数据。',
 '反馈材料不包含真实密码或令牌；下载问题需补充浏览器下载提示。']),
]:
    doc.add_heading(title,2)
    for j,text in enumerate(items,1):
        p=doc.add_paragraph(f'{j}. {text}','Body Text');p.paragraph_format.first_line_indent=Pt(28)
for title,shot,caption,text in [
 ('附录　阈值保存结果','monitor-rules-saved','图 附－1　刷新后 CPU 阈值保持为 88','对照第十一章，保存阈值后回到列表并刷新，核对 CPU > 88 %。该修改不会重新计算已有告警。'),
 ('附录　告警关闭结果','incident-closed','图 附－2　告警状态和按钮在关闭后的变化','对照第十四章，确认关闭后核对“已关闭”标签及不可用按钮。告警状态改变不表示主机负载已经恢复。'),
 ('附录　流程完成结果','automation-completed','图 附－3　确认执行后的完成提示','对照第十六章，在所选流程卡片查看完成提示，再到执行记录核对新增记录；仅打开确认窗口不会生成记录。'),
 ('附录　知识编辑填写','knowledge-new','图 附－4　标题与正文的填写范围','对照第十八章，先填写标题，再按检查对象、检查内容和核验记录组织正文；确认无误后保存，不保存则选择取消。'),
]:
    new_heading(title);para(text);figure(shot,caption,height=16)
for start in range(0,len(CHAPTERS),2):
    new_heading('附录　页面入口图示' if start==0 else f'附录　页面入口图示　{start//2+1}')
    for offset,c in enumerate(CHAPTERS[start:start+2],start+1):
        figure(c['shot'],f'图 附－{offset+4}　{c["captions"][0]}',height=8.8)
args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output);print(args.output)
